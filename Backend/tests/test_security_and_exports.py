import re
from io import BytesIO
from datetime import date
from decimal import Decimal
from types import SimpleNamespace
from uuid import UUID

import pytest
from fastapi.testclient import TestClient
from fastapi import HTTPException
from openpyxl import load_workbook

from core.database import get_db
from gateway.middlewares.auth_middleware import get_current_admin
from main import app
from modules.analitica.controller import AnaliticaController
from modules.ingesta.controller import IngestaController
from modules.transformacion.controller import transformacionController
from modules.auth.service.AuthService import AuthService
from shared.errors import AuthError
from modules.analitica.services.AnaliticaService import (
    FalloEjecucionAnaliticaError,
    registrar_ejecucion_analitica,
)
from shared.exporting import render_export


client = TestClient(app)


def test_jwt_hs256_round_trip_and_invalid_token_rejection():
    service = AuthService.__new__(AuthService)
    token, jti, _expires_at = service.create_access_token(7, "admin-test")

    payload = service.decode_token(token)
    assert payload["sub"] == "7"
    assert payload["jti"] == jti

    with pytest.raises(AuthError) as exc_info:
        service.decode_token("malformed-token")
    assert exc_info.value.code == "expired_token"


def test_admin_operations_reject_anonymous_requests():
    attempts = [
        ("post", "/api/auth/register", {"username": "nuevo", "email": "a@example.com", "password": "long-password-123"}),
        ("post", "/api/procesados/reprocesar", {"forzar_reproceso": False}),
        ("post", "/api/analitica/riesgo/calcular", None),
        ("get", "/api/analitica/outliers", None),
        ("get", "/api/analitica/duplicados/resumen", None),
        ("get", "/api/analitica/directas", None),
        ("get", "/api/analitica/riesgo", None),
        ("get", "/api/analitica/pesos", None),
        ("get", "/api/ingesta/fuentes/", None),
        ("get", "/api/jobs/resumen", None),
        ("get", "/api/procesados/logs/export/csv", None),
    ]
    for method, path, body in attempts:
        response = client.request(method, path, json=body)
        assert response.status_code in (401, 403), (path, response.text)


def test_openapi_requires_authentication_for_every_write_and_admin_read():
    protected_read_prefixes = (
        "/api/ingesta/fuentes",
        "/api/procesados/logs",
        "/api/analitica",
        "/api/jobs",
    )
    public_post_paths = {
        "/api/auth/login",
        "/api/procesados/export/{formato}/jobs",
    }
    for path, operations in app.openapi()["paths"].items():
        for method, operation in operations.items():
            if method not in {"get", "post", "put", "patch", "delete"}:
                continue
            must_be_protected = (
                (method != "get" and path not in public_post_paths)
                or path.startswith(protected_read_prefixes)
            )
            if must_be_protected:
                assert operation.get("security"), f"Missing auth: {method.upper()} {path}"


def test_every_bearer_secured_operation_rejects_anonymous_requests_before_database(monkeypatch):
    def database_must_not_be_reached():
        raise AssertionError("An anonymous request reached the database dependency")

    monkeypatch.setitem(app.dependency_overrides, get_db, database_must_not_be_reached)
    checked = 0
    for path, operations in app.openapi()["paths"].items():
        for method, operation in operations.items():
            if method not in {"get", "post", "put", "patch", "delete"}:
                continue
            if not any(
                "HTTPBearer" in requirement
                for requirement in operation.get("security", [])
            ):
                continue

            concrete_path = re.sub(r"\{[^/]+\}", "1", path)
            kwargs = {"json": {}} if method in {"post", "put", "patch"} else {}
            response = client.request(method, concrete_path, **kwargs)
            assert response.status_code in (401, 403), (
                f"Anonymous {method.upper()} {path} returned "
                f"{response.status_code}: {response.text}"
            )
            checked += 1

    assert checked > 0


def test_malformed_bearer_token_is_rejected_as_unauthorized():
    response = client.get("/api/auth/me", headers={"Authorization": "Bearer malformed"})
    assert response.status_code == 401


def test_contract_search_rejects_unknown_query_parameters_before_database_access():
    app.dependency_overrides[get_db] = lambda: object()
    try:
        response = client.get("/api/procesados/search?limit=10&limti=25")
    finally:
        app.dependency_overrides.pop(get_db, None)

    assert response.status_code == 422
    assert "limti" in response.json()["detail"]


@pytest.mark.parametrize(
    "path",
    [
        "/api/procesados/?page=1&pagina=3",
        "/api/procesados/autocomplete?q=acme&size=10",
        "/api/procesados/metricas/dashboard?limit=10",
        "/api/procesados/metricas/top-proveedores?limite=10",
        "/api/procesados/sospechosos?page=1&limit=20",
        "/api/procesados/anomalias/?motivo=INVALIDO",
        "/api/procesados/estadisticas/campos-faltantes?x=1",
        "/api/procesados/metricas/calidad?x=1",
        "/api/procesados/metricas/campos-faltantes?x=1",
        "/api/procesados/1?fields=foo",
    ],
)
def test_transformation_query_validation_happens_before_repository_access(monkeypatch, path):
    class Repository:
        def __init__(self, _db):
            raise AssertionError("La validación de consulta debe ocurrir antes del repositorio")

    monkeypatch.setattr(transformacionController, "TransformacionRepository", Repository)
    app.dependency_overrides[get_db] = lambda: object()
    try:
        response = client.get(path)
    finally:
        app.dependency_overrides.pop(get_db, None)

    assert response.status_code == 422, (path, response.text)


@pytest.mark.parametrize(
    ("path", "expected_status"),
    [
        ("/api/procesados/search?fecha_inicio=2026-12-01&fecha_fin=2026-01-01", 400),
        ("/api/procesados/search?valor_min=20&valor_max=10", 400),
        ("/api/procesados/search?nivel_confianza_min=90&nivel_confianza_max=10", 400),
        ("/api/procesados/search?valor_min=-1", 422),
        ("/api/procesados/search?nivel_confianza_max=101", 422),
    ],
)
def test_contract_search_rejects_invalid_ranges_before_repository_access(monkeypatch, path, expected_status):
    class Repository:
        def __init__(self, _db):
            raise AssertionError("La validación debe ocurrir antes del repositorio")

    monkeypatch.setattr(transformacionController, "TransformacionRepository", Repository)
    app.dependency_overrides[get_db] = lambda: object()
    try:
        response = client.get(path)
    finally:
        app.dependency_overrides.pop(get_db, None)

    assert response.status_code == expected_status, (path, response.text)


@pytest.mark.parametrize(
    "path",
    [
        "/api/analitica/outliers?page=1&page_szie=20",
        "/api/analitica/duplicados/resumen?run=latest",
        "/api/analitica/riesgo?risk=ALTO",
        "/api/analitica/outliers?run_id=not-a-uuid",
        "/api/analitica/duplicados?riesgo=CRITICO",
        "/api/analitica/directas?score_minimo=90&score_maximo=10",
        "/api/analitica/directas?porcentaje_minimo=90&porcentaje_maximo=10",
        "/api/analitica/outliers?score_minimo=inf",
        "/api/analitica/duplicados?score_minimo=nan",
        "/api/analitica/directas?score_maximo=inf",
        "/api/analitica/riesgo?score_minimo=-inf",
    ],
)
def test_analytics_query_validation_happens_before_service_access(monkeypatch, path):
    class Service:
        def __init__(self, _db):
            raise AssertionError("La validación de consulta debe ocurrir antes del servicio")

    monkeypatch.setattr(AnaliticaController, "AnaliticaService", Service)
    app.dependency_overrides[get_db] = lambda: object()
    app.dependency_overrides[get_current_admin] = lambda: SimpleNamespace(id=1)
    try:
        response = client.get(path)
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_admin, None)

    assert response.status_code == 422, (path, response.text)


@pytest.mark.parametrize(
    ("path", "body"),
    [
        (
            "/api/analitica/outliers/calcular",
            {"fecha_desde": "2026-02-01", "fecha_hasta": "2026-01-01"},
        ),
        (
            "/api/analitica/duplicados/calcular",
            {"fecha_desde": "2026-02-01", "fecha_hasta": "2026-01-01"},
        ),
        (
            "/api/analitica/directas/calcular",
            {"fecha_desde": "2026-02-01", "fecha_hasta": "2026-01-01"},
        ),
        ("/api/analitica/pesos/OUTLIER", {"peso": "1000.00"}),
    ],
)
def test_analytics_ranges_and_weight_bounds_fail_before_service_access(monkeypatch, path, body):
    class Service:
        def __init__(self, _db):
            raise AssertionError("La validación debe ocurrir antes del servicio")

    monkeypatch.setattr(AnaliticaController, "AnaliticaService", Service)
    app.dependency_overrides[get_current_admin] = lambda: object()
    app.dependency_overrides[get_db] = lambda: object()
    try:
        response = client.request("PUT" if path.endswith("/OUTLIER") else "POST", path, json=body)
    finally:
        app.dependency_overrides.pop(get_current_admin, None)
        app.dependency_overrides.pop(get_db, None)

    assert response.status_code == 422, (path, response.text)


def test_ingestion_query_validation_happens_before_service_access(monkeypatch):
    app.dependency_overrides[get_current_admin] = lambda: object()
    app.dependency_overrides[IngestaController.get_service] = lambda: object()
    try:
        response = client.get(
            "/api/ingesta/fuentes/sincronizaciones/pagina?page=1&size=20&pagesize=20"
        )
    finally:
        app.dependency_overrides.pop(get_current_admin, None)
        app.dependency_overrides.pop(IngestaController.get_service, None)

    assert response.status_code == 422
    assert "pagesize" in response.json()["detail"]


def test_ingestion_export_rejects_invalid_format_before_repository_access(monkeypatch):
    class Repository:
        def __init__(self, _db):
            raise AssertionError("El formato debe validarse antes del repositorio")

    monkeypatch.setattr(IngestaController, "IngestaRepository", Repository)
    app.dependency_overrides[get_current_admin] = lambda: object()
    app.dependency_overrides[get_db] = lambda: object()
    try:
        response = client.get("/api/ingesta/fuentes/sincronizaciones/export/xml")
    finally:
        app.dependency_overrides.pop(get_current_admin, None)
        app.dependency_overrides.pop(get_db, None)

    assert response.status_code == 422


def test_job_query_validation_happens_before_database_query():
    app.dependency_overrides[get_current_admin] = lambda: object()
    app.dependency_overrides[get_db] = lambda: object()
    try:
        response = client.get(
            "/api/jobs/550e8400-e29b-41d4-a716-446655440000?include=events"
        )
    finally:
        app.dependency_overrides.pop(get_current_admin, None)
        app.dependency_overrides.pop(get_db, None)

    assert response.status_code == 422
    assert "include" in response.json()["detail"]


def test_export_neutralizes_spreadsheet_formulas():
    rows = [["=HYPERLINK(\"https://example.com\")", "  +SUM(1,1)", 42, True, -7]]
    columns = ["A", "B", "Count", "Active", "Balance"]
    csv_response = render_export(columns, rows, "csv", "test", "Test")
    assert b"'=HYPERLINK" in csv_response.body
    assert b"'  +SUM" in csv_response.body
    assert b";42;True;-7" in csv_response.body

    xlsx_response = render_export(columns, rows, "xlsx", "test", "Test")
    workbook = load_workbook(BytesIO(xlsx_response.body), read_only=True)
    sheet = workbook.active
    cells = list(sheet.values)[1]
    assert cells[0].startswith("'=")
    assert cells[1].startswith("'  +")
    assert not cells[0].startswith("=")
    assert cells[2:] == (42, True, -7)
    assert sheet["C2"].data_type == "n"
    assert sheet["D2"].data_type == "b"
    assert sheet["E2"].data_type == "n"
    assert sheet["A2"].data_type == "s"
    assert sheet["B2"].data_type == "s"
    workbook.close()


def test_pdf_export_is_valid():
    response = render_export(["ID", "Nombre"], [[1, "Contrato"]], "pdf", "test", "Test")
    assert response.body.startswith(b"%PDF")


def test_contract_export_applies_filters_and_exports_all_filtered_rows(monkeypatch):
    calls = []

    class Repository:
        def __init__(self, _db):
            pass

        def search_contratos(self, filters, *, skip, limit, sort, order):
            calls.append((filters, skip, limit, sort, order))
            rows = [
                SimpleNamespace(
                    id=index,
                    entidad_normalizada="ENTIDAD",
                    proveedor_normalizado="=2+2" if index == 1 else "PROVEEDOR",
                    modalidad_contratacion="MODALIDAD",
                    valor_total_normalizado=Decimal("123.00"),
                    fecha_publicacion_normalizada=date(2026, 1, 1),
                    estado_normalizado="PUBLICADO",
                    nivel_confianza=80,
                    es_incompleto=False,
                    es_sospechoso=True,
                    clasificacion_riesgo="SIN_EVALUAR",
                )
                for index in range(1, 43)
            ]
            return rows, len(rows)

    monkeypatch.setattr(transformacionController, "TransformacionRepository", Repository)
    app.dependency_overrides[get_db] = lambda: object()
    try:
        response = client.get(
            "/api/procesados/export/xlsx?limit=10&offset=20&q=acme&solo_sospechosos=true&order=asc"
        )
    finally:
        app.dependency_overrides.pop(get_db, None)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    sheet = workbook.active
    exported = list(sheet.values)
    assert len(exported) == 43
    assert exported[1][2] == "'=2+2"
    workbook.close()

    filters, skip, limit, sort, order = calls[0]
    assert filters.query == "acme"
    assert filters.solo_sospechosos is True
    assert skip == 0
    assert limit == 10000
    assert sort is None
    assert order == "asc"


def test_contract_export_rejects_invalid_format_before_query(monkeypatch):
    class Repository:
        def __init__(self, _db):
            raise AssertionError("No se debe consultar antes de validar el formato")

    monkeypatch.setattr(transformacionController, "TransformacionRepository", Repository)
    app.dependency_overrides[get_db] = lambda: object()
    try:
        response = client.get("/api/procesados/export/xml")
    finally:
        app.dependency_overrides.pop(get_db, None)

    assert response.status_code == 422


def test_analytics_calculation_enqueues_instead_of_running_inside_api(monkeypatch):
    queued = {}

    def fake_enqueue(_db, *, kind, payload, resource_key):
        queued.update(kind=kind, payload=payload, resource_key=resource_key)
        return SimpleNamespace(
            public_id=UUID(int=1), kind=kind, status="PENDIENTE"
        ), True

    monkeypatch.setattr(AnaliticaController, "enqueue_job", fake_enqueue)
    app.dependency_overrides[get_db] = lambda: object()
    app.dependency_overrides[get_current_admin] = lambda: object()
    try:
        response = client.post("/api/analitica/outliers/calcular", json={})
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_admin, None)

    assert response.status_code == 202
    assert response.json()["id"] == str(UUID(int=1))
    assert response.json()["kind"] == "ANALITICA"
    assert queued == {
        "kind": "ANALITICA",
        "payload": {
            "tipo": "OUTLIERS",
            "parametros": {
                "campo": "valor_total_normalizado",
                "fecha_campo": None,
                "fecha_desde": None,
                "fecha_hasta": None,
                "modalidad": None,
            },
        },
        "resource_key": "analitica:OUTLIERS",
    }


def test_analytics_failure_reference_matches_recorded_run_and_server_log(caplog):
    secret_detail = "internal-driver-password=must-not-leak"

    class Repository:
        error_id = None

        def obtener_firma_universo(self):
            return {"total": 3}

        def iniciar_ejecucion_analitica(self, **_kwargs):
            pass

        def finalizar_ejecucion_analitica(self, _run_id, **kwargs):
            self.error_id = kwargs["error_id"]

    repository = Repository()

    class Service:
        db = SimpleNamespace(rollback=lambda: None)
        repo = repository

    @registrar_ejecucion_analitica("RIESGO", recibe_request=False)
    def fail_run(self, *, _run_id=None):
        raise ValueError(secret_detail)

    caplog.set_level("ERROR")
    with pytest.raises(FalloEjecucionAnaliticaError) as exc_info:
        fail_run(Service())

    assert repository.error_id == UUID(exc_info.value.error_id)
    assert exc_info.value.error_id in caplog.text
    assert secret_detail not in caplog.text


def test_scheduler_does_not_log_exception_details(monkeypatch, caplog):
    from core import scheduler as scheduler_module

    secret_detail = "database-password=must-not-leak"
    source = SimpleNamespace(id=71, nombre="SECOP", ultima_sync=None, frecuencia_dias=15)

    class Repository:
        def __init__(self, _db):
            pass

        def get_activas(self):
            return [source]

    class Database:
        closed = False

        def close(self):
            self.closed = True

    database = Database()

    def fail_enqueue(*_args, **_kwargs):
        raise RuntimeError(secret_detail)

    monkeypatch.setattr(scheduler_module, "SessionLocal", lambda: database)
    monkeypatch.setattr(scheduler_module, "IngestaRepository", Repository)
    monkeypatch.setattr(scheduler_module, "enqueue_job", fail_enqueue)
    caplog.set_level("ERROR")

    scheduler_module.sincronizar_fuentes_activas()

    assert database.closed
    assert "fuente_id=71" in caplog.text
    assert "RuntimeError" in caplog.text
    assert "referencia=" in caplog.text
    assert secret_detail not in caplog.text
